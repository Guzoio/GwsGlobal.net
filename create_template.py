"""
Script auxiliar para gerar automaticamente o arquivo 'template_proposta.docx'
com o papel timbrado, formatação institucional e todas as tags do docxtpl.
"""

import os
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def criar_template():
    doc = Document()

    # Margens padrão A4 (2 cm)
    for section in doc.sections:
        section.top_margin = Mm(20)
        section.bottom_margin = Mm(20)
        section.left_margin = Mm(20)
        section.right_margin = Mm(20)

    # 1. Cabeçalho / Papel Timbrado Estático
    p_header = doc.add_paragraph()
    p_header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p_header.add_run("COMERCIAL DISTRIBUIDORA BRASIL LTDA.\n")
    r1.bold = True
    r1.font.size = Pt(13)
    r1.font.color.rgb = RGBColor(15, 44, 89)

    r2 = p_header.add_run(
        "CNPJ: 12.345.678/0001-90 | Inscrição Estadual: 123.456.789.000\n"
        "Av. Paulista, 1000 - Cj. 101 - São Paulo/SP | CEP: 01310-100\n"
        "Telefone: (11) 3456-7890 | E-mail: licitacoes@empresa.com.br"
    )
    r2.font.size = Pt(8.5)
    r2.font.color.rgb = RGBColor(100, 116, 139)

    # Divisória
    p_line = doc.add_paragraph("―" * 55)
    p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_line.runs[0].font.color.rgb = RGBColor(203, 213, 225)

    # Título
    p_tit = doc.add_paragraph()
    p_tit.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_tit = p_tit.add_run("PROPOSTA COMERCIAL DE PREÇOS")
    r_tit.bold = True
    r_tit.font.size = Pt(14)
    r_tit.font.color.rgb = RGBColor(15, 44, 89)

    # Metadados com tags Jinja2
    p_meta = doc.add_paragraph()
    p_meta.add_run("Ao Órgão Licitante: ").bold = True
    p_meta.add_run("{{ orgao }}\n")
    p_meta.add_run("Processo / Pregão: ").bold = True
    p_meta.add_run("{{ processo_pregao }}    |    ")
    p_meta.add_run("Modalidade: ").bold = True
    p_meta.add_run("{{ modalidade }}\n")
    p_meta.add_run("Data de Emissão: ").bold = True
    p_meta.add_run("{{ data_hoje }}")

    doc.add_paragraph()

    # Tabela Comercial
    table = doc.add_table(rows=2, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Item", "Descrição", "Marca/Modelo", "Qtd.", "Valor Unit.", "Valor Total"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].bold = True
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
        shd = parse_xml(r'<w:shd {} w:fill="0F2C59"/>'.format(nsdecls('w')))
        cell._tc.get_or_add_tcPr().append(shd)

    # Linha dinâmica Jinja2 para docxtpl
    row_tags = [
        "{% tr for item in itens %}{{ item.num_item }}",
        "{{ item.descricao_curta }}",
        "{{ item.marca }}",
        "{{ item.quantidade }}",
        "{{ item.valor_unitario }}",
        "{{ item.valor_total }}{% tr endfor %}"
    ]
    for i, t in enumerate(row_tags):
        table.cell(1, i).text = t

    doc.add_paragraph()

    # Total e Extenso
    p_tot = doc.add_paragraph()
    p_tot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_t_lbl = p_tot.add_run("TOTAL GERAL: ")
    r_t_lbl.bold = True
    p_tot.add_run("{{ total_geral }}")

    p_ext = doc.add_paragraph()
    p_ext.add_run("Valor por Extenso: ").bold = True
    p_ext.add_run("{{ total_geral }} ({{ extenso }})")

    # Assinatura
    p_ass = doc.add_paragraph("\n\n")
    p_ass.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ass.add_run("_____________________________________________\n").bold = True
    p_ass.add_run("COMERCIAL DISTRIBUIDORA BRASIL LTDA.\n").bold = True
    p_ass.add_run("Responsável Legal / Licitações")

    # Catálogo
    doc.add_page_break()
    p_cat = doc.add_paragraph()
    p_cat.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_cat = p_cat.add_run("ANEXO - CATÁLOGO E ESPECIFICAÇÕES TÉCNICAS DOS PRODUTOS")
    r_cat.bold = True
    r_cat.font.size = Pt(13)
    r_cat.font.color.rgb = RGBColor(15, 44, 89)

    doc.add_paragraph("Detalhamento técnico e documentação fotográfica dos produtos cotados:\n")

    cat_table = doc.add_table(rows=1, cols=2)
    cat_table.cell(0, 0).text = "{% for item in itens %}{{ item.imagem_docxtpl }}"
    cat_table.cell(0, 1).text = (
        "ITEM {{ item.num_item }} — {{ item.descricao_curta }}\n"
        "Marca: {{ item.marca }}\n\n"
        "Especificações Técnicas:\n"
        "{{ item.descricao_tecnica }}{% endfor %}"
    )

    doc.save("template_proposta.docx")
    print("Arquivo 'template_proposta.docx' criado com sucesso!")

if __name__ == "__main__":
    criar_template()
